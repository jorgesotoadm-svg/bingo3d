import io
import os
from fpdf import FPDF
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd
import streamlit as st

# Configuración de la página en formato ancho
st.set_page_config(
    page_title="Cotizador Profesional - Bingo Impresión 3D",
    page_icon="🖨️",
    layout="wide",
)

# Estilos CSS avanzados con colores corporativos y fondo cálido
st.markdown(
    """
    <style>
    /* Fondo general de toda la aplicación */
    .stApp {
        background-color: #F4F1EA !important;
    }
    
    /* Estilos de títulos */
    h1, h2, h3 {
        color: #2B2D42 !important;
        font-family: 'Helvetica Neue', sans-serif;
    }
    
    /* Botones corporativos */
    .stButton>button {
        background-color: #2B2D42 !important;
        color: white !important;
        border-radius: 8px;
        border: none;
        font-weight: bold;
        padding: 0.5rem 1rem;
    }
    .stButton>button:hover {
        background-color: #C5A089 !important;
        color: #2B2D42 !important;
    }
    
    /* Tarjetas blancas con relieve y borde lateral corporativo */
    .card {
        background-color: #FFFFFF;
        padding: 24px;
        border-radius: 12px;
        box-shadow: 0 4px 12px rgba(0,0,0,0.08);
        margin-bottom: 20px;
        border-left: 6px solid #C5A089;
    }
    
    .item-card {
        background-color: #FAFAFA;
        padding: 20px;
        border-radius: 10px;
        border: 1px solid #E0DCD0;
        margin-bottom: 15px;
    }
    </style>
""",
    unsafe_allow_html=True,
)

# Encabezado con Logo y Título Corporativo
col_logo, col_title = st.columns([1, 5])
with col_logo:
  if os.path.exists("logo.png"):
    st.image("logo.png", width=130)
  elif os.path.exists("logo.png.jpg"):
    st.image("logo.png.jpg", width=130)
  else:
    st.markdown("### 🖨️ BINGO 3D")

with col_title:
  st.markdown(
      "<h1"
      ' style="margin-bottom:0; color:#2B2D42;">BINGO IMPRESIÓN 3D</h1>',
      unsafe_allow_html=True,
  )
  st.markdown(
      "<p style='color:#C5A089; font-weight:bold; font-size:1.1rem;'>Sistema"
      " Avanzado de Costos y Cotizaciones 3D</p>",
      unsafe_allow_html=True,
  )

st.markdown("---")

# 1. Datos de la Empresa Emisora
st.markdown('<div class="card">', unsafe_allow_html=True)
st.subheader("🏢 Datos de la Empresa (Emisor)")
col_e1, col_e2, col_e3, col_e4 = st.columns(4)
with col_e1:
  empresa_nombre = st.text_input(
      "Nombre Empresa", "Bingo Impresión 3D"
  )
with col_e2:
  empresa_rut = st.text_input("RUT Empresa", "76.XXX.XXX-X")
with col_e3:
  empresa_telefono = st.text_input("Teléfono Emisor", "+56 9 XXXXXXXX")
with col_e4:
  empresa_correo = st.text_input("Correo Emisor", "contacto@bingo3d.cl")
st.markdown("</div>", unsafe_allow_html=True)

# 2. Datos del Cliente y Detalles de la Cotización
col_sec1, col_sec2 = st.columns(2)

with col_sec1:
  st.markdown('<div class="card">', unsafe_allow_html=True)
  st.subheader("👤 Datos del Cliente")
  cliente_nombre = st.text_input(
      "Nombre / Empresa Cliente", "Cliente General"
  )
  cliente_rut = st.text_input("RUT Cliente", "12.345.678-9")
  cliente_tel = st.text_input("Teléfono Cliente", "+56 9 ...")
  cliente_email = st.text_input("Correo Cliente", "cliente@correo.com")
  st.markdown("</div>", unsafe_allow_html=True)

with col_sec2:
  st.markdown('<div class="card">', unsafe_allow_html=True)
  st.subheader("📄 Detalles de la Cotización")
  num_cotizacion = st.text_input("N° de Cotización", "COT-2026-001")
  fecha_cot = st.date_input("Fecha de Emisión")
  validez = st.text_input("Validez de la Oferta", "15 días")
  forma_pago = st.text_input(
      "Forma de Pago", "50% adelanto, 50% saldo contra entrega"
  )
  st.markdown("</div>", unsafe_allow_html=True)

st.markdown("---")

# 3. Calculador Avanzado de Ítems e Impresiones 3D
st.markdown('<div class="card">', unsafe_allow_html=True)
st.subheader("🛠️ Detalle de Modelos y Costos 3D")

if "lista_items" not in st.session_state:
  st.session_state["lista_items"] = [{
      "descripcion": "Figura / Modelo 3D Personalizado",
      "tipo_filamento": "PLA",
      "precio_rollo": 18000.0,
      "peso_rollo_gramos": 1000.0,
      "peso_gramos": 50.0,
      "tiempo_horas": 5.0,
      "costo_hora_maquina": 500.0,
      "mano_obra": 3000.0,
      "cantidad": 1,
  }]


def add_item():
  st.session_state["lista_items"].append({
      "descripcion": "Nuevo Modelo 3D",
      "tipo_filamento": "PLA",
      "precio_rollo": 18000.0,
      "peso_rollo_gramos": 1000.0,
      "peso_gramos": 30.0,
      "tiempo_horas": 3.0,
      "costo_hora_maquina": 500.0,
      "mano_obra": 2000.0,
      "cantidad": 1,
  })


def remove_item(index):
  st.session_state["lista_items"].pop(index)


subtotal_general = 0.0

for i, item in enumerate(st.session_state["lista_items"]):
  st.markdown(
      f'<div class="item-card"><h4 style="color:#2B2D42; margin-top:0;">Modelo'
      f" #{i+1}</h4>",
      unsafe_allow_html=True,
  )

  c1, c2, c3 = st.columns([3, 1, 2])
  with c1:
    st.session_state["lista_items"][i]["descripcion"] = st.text_input(
        f"Descripción de la Figura {i+1}",
        item["descripcion"],
        key=f"desc_{i}",
    )
  with c2:
    st.session_state["lista_items"][i]["cantidad"] = st.number_input(
        f"Cantidad", min_value=1, value=item["cantidad"], step=1, key=f"cant_{i}"
    )
  with c3:
    st.session_state["lista_items"][i]["tipo_filamento"] = st.selectbox(
        f"Tipo de Filamento",
        ["PLA", "PETG", "ABS", "TPU (Flexible)", "Silk / Especial", "Resina"],
        index=[
            "PLA",
            "PETG",
            "ABS",
            "TPU (Flexible)",
            "Silk / Especial",
            "Resina",
        ].index(item["tipo_filamento"])
        if item["tipo_filamento"]
        in ["PLA", "PETG", "ABS", "TPU (Flexible)", "Silk / Especial", "Resina"]
        else 0,
        key=f"fil_{i}",
    )

  c4, c5, c6, c7, c8, c9 = st.columns(6)
  with c4:
    st.session_state["lista_items"][i]["precio_rollo"] = st.number_input(
        "Precio Rollo ($)",
        min_value=0.0,
        value=float(item.get("precio_rollo", 18000.0)),
        step=1000.0,
        key=f"p_rollo_{i}",
    )
  with c5:
    st.session_state["lista_items"][i]["peso_rollo_gramos"] = st.number_input(
        "Peso Rollo (g)",
        min_value=1.0,
        value=float(item.get("peso_rollo_gramos", 1000.0)),
        step=100.0,
        key=f"w_rollo_{i}",
    )
  with c6:
    st.session_state["lista_items"][i]["peso_gramos"] = st.number_input(
        "Peso Figura (g)",
        min_value=0.0,
        value=float(item["peso_gramos"]),
        step=5.0,
        key=f"peso_{i}",
    )
  with c7:
    st.session_state["lista_items"][i]["tiempo_horas"] = st.number_input(
        "Tiempo (Horas)",
        min_value=0.0,
        value=float(item["tiempo_horas"]),
        step=0.5,
        key=f"time_{i}",
    )
  with c8:
    st.session_state["lista_items"][i]["costo_hora_maquina"] = st.number_input(
        "Costo Maq/hr ($)",
        min_value=0.0,
        value=float(item["costo_hora_maquina"]),
        step=100.0,
        key=f"c_m_{i}",
    )
  with c9:
    st.session_state["lista_items"][i]["mano_obra"] = st.number_input(
        "Mano de Obra ($)",
        min_value=0.0,
        value=float(item["mano_obra"]),
        step=500.0,
        key=f"m_o_{i}",
    )

  # CÁLCULO DE COSTO POR GRAMO
  p_rollo = st.session_state["lista_items"][i]["precio_rollo"]
  w_rollo = st.session_state["lista_items"][i]["peso_rollo_gramos"]
  costo_gramo = p_rollo / w_rollo if w_rollo > 0 else 0.0

  peso_fig = st.session_state["lista_items"][i]["peso_gramos"]
  horas = st.session_state["lista_items"][i]["tiempo_horas"]
  c_maq = st.session_state["lista_items"][i]["costo_hora_maquina"]
  m_obra = st.session_state["lista_items"][i]["mano_obra"]
  cant = st.session_state["lista_items"][i]["cantidad"]

  costo_material = peso_fig * costo_gramo
  costo_impresion_maq = horas * c_maq
  precio_unitario_calculado = costo_material + costo_impresion_maq + m_obra
  total_item = precio_unitario_calculado * cant

  subtotal_general += total_item

  resumen_html = f"""
    <div style="margin-top: 12px; padding: 12px; background-color: #FFFFFF; border-radius: 6px; border: 1px solid #DCD7CD;">
        📊 <b>Costo x Gramo Calculado:</b> ${costo_gramo:.1f} | 
        Material: <b>${costo_material:,.0f}</b> | 
        Máquina: <b>${costo_impresion_maq:,.0f}</b> | 
        Mano de Obra: <b>${m_obra:,.0f}</b> 
        &nbsp;&rarr;&nbsp; 
        <b>Precio Unitario: <span style="color:#C5A089; font-size:1.2rem;">${precio_unitario_calculado:,.0f} CLP</span></b> 
        (Total ítem: <b>${total_item:,.0f} CLP</b>)
    </div>
    """
  st.markdown(resumen_html, unsafe_allow_html=True)

  if len(st.session_state["lista_items"]) > 1:
    if st.button(f"🗑️ Eliminar Modelo #{i+1}", key=f"del_{i}"):
      remove_item(i)
      st.rerun()

  st.markdown("</div>", unsafe_allow_html=True)

if st.button("➕ Agregar otro modelo 3D"):
  add_item()
  st.rerun()

st.markdown("</div>", unsafe_allow_html=True)

# 4. CONFIGURACIÓN DE MERMA / RIESGO
st.markdown('<div class="card">', unsafe_allow_html=True)
st.subheader("⚠️ Factor de Riesgo y Contingencia")
porcentaje_merma = st.slider(
    "Margen por Merma o Impresiones Fallidas (%)",
    min_value=0.0,
    max_value=20.0,
    value=5.0,
    step=1.0,
    help=(
        "Porcentaje adicional para cubrir posibles fallos de impresión, reintentos"
        " o imprevistos."
    ),
)
st.markdown("</div>", unsafe_allow_html=True)

# Cálculos financieros con factor de merma e IVA 19%
monto_merma = subtotal_general * (porcentaje_merma / 100.0)
subtotal_sin_iva = subtotal_general + monto_merma
iva_porcentaje = 0.19
monto_iva = subtotal_sin_iva * iva_porcentaje
total_con_iva = subtotal_sin_iva + monto_iva

# 5. Resumen Financiero con Estilo Corporativo
st.markdown('<div class="card">', unsafe_allow_html=True)
st.subheader("💰 Resumen Financiero")
col_res1, col_res2, col_res3, col_res4 = st.columns(4)

with col_res1:
  st.metric(
      label="Subtotal Producción", value=f"${subtotal_general:,.0f} CLP"
  )
with col_res2:
  st.metric(
      label=f"Merma ({porcentaje_merma}%)", value=f"${monto_merma:,.0f} CLP"
  )
with col_res3:
  st.metric(label="IVA (19%)", value=f"${monto_iva:,.0f} CLP")
with col_res4:
  st.markdown(
      f"<h4 style='color:#C5A089; margin:0;'>TOTAL FINAL (CON IVA)</h4>"
      f"<h2 style='color:#2B2D42; margin:0;'>${total_con_iva:,.0f} CLP</h2>",
      unsafe_allow_html=True,
  )
st.markdown("</div>", unsafe_allow_html=True)

st.markdown("---")

# 6. Generador de PDF y Excel Profesional
st.markdown('<div class="card">', unsafe_allow_html=True)
st.subheader("📥 Exportar Documentos Oficiales")

col_dl1, col_dl2 = st.columns(2)


# Función para generar PDF formal
def generar_pdf():
  pdf = FPDF()
  pdf.add_page()
  pdf.set_auto_page_break(auto=True, margin=15)

  logo_path = None
  if os.path.exists("logo.png"):
    logo_path = "logo.png"
  elif os.path.exists("logo.png.jpg"):
    logo_path = "logo.png.jpg"

  if logo_path:
    try:
      pdf.image(logo_path, 10, 10, 28)
    except:
      pass

  pdf.set_font("Helvetica", "B", 15)
  pdf.set_text_color(43, 45, 66)
  pdf.cell(
      0, 8, empresa_nombre.encode("latin-1", "replace").decode("latin-1"), 0, 1, "R"
  )

  pdf.set_font("Helvetica", "", 9)
  pdf.set_text_color(100, 100, 100)
  pdf.cell(
      0,
      4,
      f"RUT: {empresa_rut} | Tel: {empresa_telefono}".encode(
          "latin-1", "replace"
      ).decode("latin-1"),
      0,
      1,
      "R",
  )
  pdf.cell(
      0,
      4,
      f"Correo: {empresa_correo}".encode("latin-1", "replace").decode("latin-1"),
      0,
      1,
      "R",
  )

  pdf.ln(8)

  pdf.set_fill_color(244, 241, 234)
  pdf.set_font("Helvetica", "B", 10)
  pdf.set_text_color(43, 45, 66)
  pdf.cell(0, 7, "  INFORMACIÓN DE LA COTIZACIÓN", 0, 1, "L", fill=True)

  pdf.set_font("Helvetica", "", 9)
  pdf.cell(
      100,
      5,
      f"N° Cotización: {num_cotizacion}".encode("latin-1", "replace").decode("latin-1"),
      0,
      0,
  )
  pdf.cell(
      0,
      5,
      f"Fecha: {str(fecha_cot)}".encode("latin-1", "replace").decode("latin-1"),
      0,
      1,
  )
  pdf.cell(
      100,
      5,
      f"Cliente: {cliente_nombre}".encode("latin-1", "replace").decode("latin-1"),
      0,
      0,
  )
  pdf.cell(
      0,
      5,
      f"Validez: {validez}".encode("latin-1", "replace").decode("latin-1"),
      0,
      1,
  )
  pdf.cell(
      100,
      5,
      f"RUT Cliente: {cliente_rut}".encode("latin-1", "replace").decode("latin-1"),
      0,
      0,
  )
  pdf.cell(
      0,
      5,
      f"Forma de Pago: {forma_pago}".encode("latin-1", "replace").decode("latin-1"),
      0,
      1,
  )

  pdf.ln(6)

  pdf.set_font("Helvetica", "B", 9)
  pdf.set_fill_color(43, 45, 66)
  pdf.set_text_color(255, 255, 255)

  col_widths = [85, 15, 25, 30, 35]
  headers = ["Descripción", "Cant.", "Filamento", "P. Unitario", "Total"]

  for i, h_text in enumerate(headers):
    pdf.cell(
        col_widths[i],
        7,
        h_text.encode("latin-1", "replace").decode("latin-1"),
        1,
        0,
        "C",
        fill=True,
    )
  pdf.ln()

  pdf.set_font("Helvetica", "", 9)
  pdf.set_text_color(0, 0, 0)

  for item in st.session_state["lista_items"]:
    p_rollo = item.get("precio_rollo", 18000.0)
    w_rollo = item.get("peso_rollo_gramos", 1000.0)
    c_g = p_rollo / w_rollo if w_rollo > 0 else 0.0
    peso_fig = item["peso_gramos"]
    h = item["tiempo_horas"]
    cm = item["costo_hora_maquina"]
    mo = item["mano_obra"]
    cant = item["cantidad"]

    p_unit = (peso_fig * c_g) + (h * cm) + mo
    t_item = p_unit * cant

    pdf.cell(
        col_widths[0],
        6,
        str(item["descripcion"])
        .encode("latin-1", "replace")
        .decode("latin-1"),
        1,
        0,
        "L",
    )
    pdf.cell(
        col_widths[1], 6, str(cant), 1, 0, "C"
    )
    pdf.cell(
        col_widths[2],
        6,
        str(item["tipo_filamento"])
        .encode("latin-1", "replace")
        .decode("latin-1"),
        1,
        0,
        "C",
    )
    pdf.cell(col_widths[3], 6, f"${p_unit:,.0f}", 1, 0, "R")
    pdf.cell(col_widths[4], 6, f"${t_item:,.0f}", 1, 1, "R")

  pdf.ln(5)

  pdf.set_font("Helvetica", "B", 9)
  pdf.set_x(110)
  pdf.cell(45, 5, "Subtotal Producción:", 0, 0, "R")
  pdf.set_font("Helvetica", "", 9)
  pdf.cell(35, 5, f"${subtotal_general:,.0f} CLP", 0, 1, "R")

  if porcentaje_merma > 0:
    pdf.set_font("Helvetica", "B", 9)
    pdf.set_x(110)
    pdf.cell(45, 5, f"Merma / Contingencia ({porcentaje_merma}%):", 0, 0, "R")
    pdf.set_font("Helvetica", "", 9)
    pdf.cell(35, 5, f"${monto_merma:,.0f} CLP", 0, 1, "R")

  pdf.set_font("Helvetica", "B", 9)
  pdf.set_x(110)
  pdf.cell(45, 5, "Subtotal Neto:", 0, 0, "R")
  pdf.set_font("Helvetica", "", 9)
  pdf.cell(35, 5, f"${subtotal_sin_iva:,.0f} CLP", 0, 1, "R")

  pdf.set_font("Helvetica", "B", 9)
  pdf.set_x(110)
  pdf.cell(45, 5, "IVA (19%):", 0, 0, "R")
  pdf.set_font("Helvetica", "", 9)
  pdf.cell(35, 5, f"${monto_iva:,.0f} CLP", 0, 1, "R")

  pdf.set_fill_color(197, 160, 137)
  pdf.set_text_color(255, 255, 255)
  pdf.set_font("Helvetica", "B", 10)
  pdf.set_x(110)
  pdf.cell(45, 7, "TOTAL FINAL:", 0, 0, "R", fill=True)
  pdf.cell(35, 7, f"${total_con_iva:,.0f} CLP", 0, 1, "R", fill=True)

  return bytes(pdf.output())


# Función para generar Excel profesional con logo y formato tipo
def generar_excel():
  wb = openpyxl.Workbook()
  ws = wb.active
  ws.title = "Cotización"
  ws.views.sheetView[0].showGridLines = True

  charcoal = "2B2D42"
  bronze = "C5A089"
  light_gray = "FAFAFA"
  white = "FFFFFF"

  font_title = Font(name="Helvetica", size=14, bold=True, color=charcoal)
  font_subtitle = Font(name="Helvetica", size=10, bold=True, color=bronze)
  font_header = Font(name="Helvetica", size=10, bold=True, color=white)
  font_bold = Font(name="Helvetica", size=10, bold=True, color="000000")
  font_normal = Font(name="Helvetica", size=10, color="000000")

  fill_header = PatternFill(
      start_color=charcoal, end_color=charcoal, fill_type="solid"
  )
  fill_total = PatternFill(
      start_color=bronze, end_color=bronze, fill_type="solid"
  )
  fill_zebra = PatternFill(
      start_color=light_gray, end_color=light_gray, fill_type="solid"
  )

  thin_border = Border(
      left=Side(style="thin", color="DCDCDC"),
      right=Side(style="thin", color="DCDCDC"),
      top=Side(style="thin", color="DCDCDC"),
      bottom=Side(style="thin", color="DCDCDC"),
  )

  logo_path = None
  if os.path.exists("logo.png"):
    logo_path = "logo.png"
  elif os.path.exists("logo.png.jpg"):
    logo_path = "logo.png.jpg"

  if logo_path:
    try:
      img = openpyxl.drawing.image.Image(logo_path)
      img.width = 110
      img.height = 55
      ws.add_image(img, "A1")
    except Exception:
      pass

  ws["C1"] = empresa_nombre.upper()
  ws["C1"].font = font_title
  ws["C2"] = "SISTEMA DE COTIZACIÓN - IMPRESIÓN 3D"
  ws["C2"].font = font_subtitle

  ws["I1"] = "N° Cotización:"
  ws["I1"].font = font_bold
  ws["J1"] = num_cotizacion
  ws["I2"] = "Fecha Emisión:"
  ws["I2"].font = font_bold
  ws["J2"] = str(fecha_cot)
  ws["I3"] = "Validez:"
  ws["I3"].font = font_bold
  ws["J3"] = validez
  ws["I4"] = "Forma de Pago:"
  ws["I4"].font = font_bold
  ws["J4"] = forma_pago

  for row in range(1, 5):
    ws[f"J{row}"].font = font_normal

  ws["A5"] = "DATOS DEL EMISOR"
  ws["A5"].font = font_bold
  ws["A6"] = f"RUT: {empresa_rut}"
  ws["A7"] = f"Teléfono: {empresa_telefono}"
  ws["A8"] = f"Correo: {empresa_correo}"

  ws["F5"] = "DATOS DEL CLIENTE"
  ws["F5"].font = font_bold
  ws["F6"] = f"Cliente: {cliente_nombre}"
  ws["F7"] = f"RUT: {cliente_rut}"
  ws["F8"] = f"Teléfono: {cliente_tel}"
  ws["F9"] = f"Correo: {cliente_email}"

  start_row = 12
  ws.cell(
      row=start_row, column=1, value="DETALLE DE MODELOS Y COSTOS 3D"
  ).font = font_title

  headers = [
      "Ítem",
      "Descripción",
      "Filamento",
      "Cant.",
      "Precio Rollo",
      "Peso Rollo (g)",
      "Peso Fig. (g)",
      "Costo Mat/g",
      "Horas Imp.",
      "Costo Maq/hr",
      "Mano de Obra",
      "Precio Unit. (CLP)",
      "Total Parcial (CLP)",
  ]
  header_row = start_row + 2

  for col_idx, header in enumerate(headers, 1):
    cell = ws.cell(row=header_row, column=col_idx, value=header)
    cell.font = font_header
    cell.fill = fill_header
    cell.alignment = Alignment(
        horizontal="center", vertical="center", wrap_text=True
    )

  current_row = header_row + 1
  for idx, item in enumerate(st.session_state["lista_items"], 1):
    p_rollo = item.get("precio_rollo", 18000.0)
    w_rollo = item.get("peso_rollo_gramos", 1000.0)
    c_g = p_rollo / w_rollo if w_rollo > 0 else 0.0

    peso_fig = item["peso_gramos"]
    h = item["tiempo_horas"]
    cm = item["costo_hora_maquina"]
    mo = item["mano_obra"]
    cant = item["cantidad"]

    c_mat = peso_fig * c_g
    c_maq = h * cm
    p_unit = c_mat + c_maq + mo
    t_item = p_unit * cant

    row_data = [
        idx,
        item["descripcion"],
        item["tipo_filamento"],
        cant,
        p_rollo,
        w_rollo,
        peso_fig,
        round(c_g, 2),
        h,
        cm,
        mo,
        p_unit,
        t_item,
    ]

    for col_idx, val in enumerate(row_data, 1):
      cell = ws.cell(row=current_row, column=col_idx, value=val)
      cell.font = font_normal
      cell.border = thin_border
      if col_idx in [4, 5, 6, 7, 8, 9, 10, 11, 12, 13]:
        cell.alignment = Alignment(horizontal="right")
        if col_idx in [5, 10, 11, 12, 13]:
          cell.number_format = '"$"#,##0'
        elif col_idx == 8:
          cell.number_format = '"$"#,##0.00'
      else:
        cell.alignment = Alignment(horizontal="left")

      if current_row % 2 == 1:
        cell.fill = fill_zebra

    current_row += 1

  current_row += 1
  totals = [
      ("Subtotal Producción (CLP)", subtotal_general),
      (f"Factor Merma / Riesgo ({porcentaje_merma}%) (CLP)", monto_merma),
      ("Subtotal Neto Ajustado (CLP)", subtotal_sin_iva),
      ("IVA (19%) (CLP)", monto_iva),
      ("TOTAL FINAL CON IVA (CLP)", total_con_iva),
  ]

  for label, val in totals:
    ws.cell(row=current_row, column=11, value=label).font = font_bold
    ws.cell(row=current_row, column=11).alignment = Alignment(
        horizontal="right"
    )
    val_cell = ws.cell(row=current_row, column=13, value=val)
    val_cell.font = font_bold
    val_cell.number_format = '"$"#,##0'
    val_cell.border = thin_border
    val_cell.alignment = Alignment(horizontal="right")

    if label.startswith("TOTAL"):
      val_cell.fill = fill_total
      val_cell.font = Font(name="Helvetica", size=11, bold=True, color=white)
    current_row += 1

  for col in ws.columns:
    max_len = max(len(str(cell.value or "")) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

  output = io.BytesIO()
  wb.save(output)
  output.seek(0)
  return output


with col_dl1:
  try:
    pdf_file = generar_pdf()
    st.download_button(
        label="📄 Descargar Cotización en Formato PDF",
        data=pdf_file,
        file_name=f"Cotizacion_{num_cotizacion.replace('/', '-')}.pdf",
        mime="application/pdf",
    )
  except Exception as e:
    st.error(f"Error al generar PDF: {e}")

with col_dl2:
  try:
    excel_file = generar_excel()
    st.download_button(
        label="📊 Descargar Planilla Excel Profesional",
        data=excel_file,
        file_name=f"Cotizacion_{num_cotizacion.replace('/', '-')}.xlsx",
        mime=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ),
    )
  except Exception as e:
    st.warning("Generando planilla Excel...")

st.markdown("</div>", unsafe_allow_html=True)